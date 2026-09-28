from abc import ABC, abstractmethod
from typing import Optional, List, Dict, Any, BinaryIO
from dataclasses import dataclass
import base64
import io
import logging

logger = logging.getLogger(__name__)


@dataclass
class ParsedContent:
    text: str
    structure: Dict[str, Any]
    metadata: Dict[str, Any]
    source_type: str


class DocumentParser(ABC):
    @abstractmethod
    async def parse(self, content: bytes, filename: str = "") -> ParsedContent:
        pass
    
    @abstractmethod
    def supports_mime_type(self, mime_type: str) -> bool:
        pass
    
    @abstractmethod
    def supports_extension(self, extension: str) -> bool:
        pass


class TextParser(DocumentParser):
    def supports_mime_type(self, mime_type: str) -> bool:
        return mime_type.startswith("text/")
    
    def supports_extension(self, extension: str) -> bool:
        return extension.lower() in [".txt", ".md", ".rst", ".log"]
    
    async def parse(self, content: bytes, filename: str = "") -> ParsedContent:
        text = content.decode("utf-8", errors="replace")
        return ParsedContent(
            text=text,
            structure={"paragraphs": text.split("\n\n")},
            metadata={"char_count": len(text), "line_count": len(text.split("\n"))},
            source_type="text"
        )


class PDFParser(DocumentParser):
    def __init__(self):
        self._pdfplumber = None
        self._pymupdf = None
    
    def _get_pdfplumber(self):
        if self._pdfplumber is None:
            try:
                import pdfplumber
                self._pdfplumber = pdfplumber
            except ImportError:
                # pypdf is a required backend dependency; pdfplumber is optional.
                self._pdfplumber = False
        return self._pdfplumber

    @staticmethod
    def extract_text(content: bytes) -> str:
        """Extract all available text without requiring optional pdfplumber."""
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(content))
        return "\n\n".join(
            f"--- PAGE {index + 1} ---\n{page.extract_text() or ''}"
            for index, page in enumerate(reader.pages)
        ).strip()
    
    def _get_pymupdf(self):
        if self._pymupdf is None:
            import fitz
            self._pymupdf = fitz
        return self._pymupdf
    
    def supports_mime_type(self, mime_type: str) -> bool:
        return mime_type == "application/pdf"
    
    def supports_extension(self, extension: str) -> bool:
        return extension.lower() == ".pdf"
    
    async def parse(self, content: bytes, filename: str = "") -> ParsedContent:
        pdfplumber = self._get_pdfplumber()

        if not pdfplumber:
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(content))
            metadata = reader.metadata or {}
            full_text = "\n\n".join(
                f"--- PAGE {index + 1} ---\n{page.extract_text() or ''}"
                for index, page in enumerate(reader.pages)
            ).strip()
            return ParsedContent(
                text=full_text,
                structure={
                    "pages": len(reader.pages),
                    "tables": [],
                    "paragraphs": full_text.split("\n\n"),
                },
                metadata={
                    "page_count": len(reader.pages),
                    "title": metadata.get("/Title"),
                    "author": metadata.get("/Author"),
                    "creator": metadata.get("/Creator"),
                },
                source_type="pdf",
            )
        
        text_parts = []
        tables = []
        images_info = []
        
        with pdfplumber.open(io.BytesIO(content)) as pdf:
            metadata = {
                "page_count": len(pdf.pages),
                "title": pdf.metadata.get("Title") if pdf.metadata else None,
                "author": pdf.metadata.get("Author") if pdf.metadata else None,
                "creator": pdf.metadata.get("Creator") if pdf.metadata else None,
            }
            
            for i, page in enumerate(pdf.pages):
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(f"--- PAGE {i+1} ---\n{page_text}")
                
                page_tables = page.extract_tables()
                if page_tables:
                    for j, table in enumerate(page_tables):
                        tables.append({
                            "page": i + 1,
                            "table_index": j,
                            "data": table
                        })
        
        full_text = "\n\n".join(text_parts)
        
        return ParsedContent(
            text=full_text,
            structure={
                "pages": len(pdf.pages) if 'pdf' in locals() else 0,
                "tables": tables,
                "paragraphs": full_text.split("\n\n")
            },
            metadata=metadata,
            source_type="pdf"
        )


class WordParser(DocumentParser):
    def __init__(self):
        self._docx = None
    
    def _get_docx(self):
        if self._docx is None:
            from docx import Document
            self._docx = Document
        return self._docx
    
    def supports_mime_type(self, mime_type: str) -> bool:
        return mime_type in [
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "application/msword"
        ]
    
    def supports_extension(self, extension: str) -> bool:
        return extension.lower() in [".docx", ".doc"]
    
    async def parse(self, content: bytes, filename: str = "") -> ParsedContent:
        Document = self._get_docx()
        doc = Document(io.BytesIO(content))
        
        text_parts = []
        headings = []
        tables = []
        images = []
        
        for i, para in enumerate(doc.paragraphs):
            if para.text.strip():
                text_parts.append(para.text)
                if para.style.name.startswith("Heading"):
                    headings.append({
                        "level": int(para.style.name.replace("Heading ", "")) if para.style.name != "Heading" else 1,
                        "text": para.text,
                        "position": i
                    })
        
        for i, table in enumerate(doc.tables):
            table_data = []
            for row in table.rows:
                table_data.append([cell.text for cell in row.cells])
            tables.append({
                "table_index": i,
                "data": table_data
            })
        
        for rel in doc.part.rels.values():
            if "image" in rel.target_ref:
                images.append({
                    "rel_id": rel.rId,
                    "content_type": rel.target_part.content_type if hasattr(rel.target_part, 'content_type') else None
                })
        
        full_text = "\n\n".join(text_parts)
        
        return ParsedContent(
            text=full_text,
            structure={
                "headings": headings,
                "tables": tables,
                "images": images,
                "paragraphs": text_parts
            },
            metadata={
                "paragraph_count": len(doc.paragraphs),
                "table_count": len(doc.tables),
                "image_count": len(images)
            },
            source_type="word"
        )


class GoogleDocsParser(DocumentParser):
    def __init__(self, credentials=None):
        self._service = None
        self._credentials = credentials
    
    def _get_service(self):
        if self._service is None:
            from googleapiclient.discovery import build
            from google.oauth2.credentials import Credentials
            
            if self._credentials:
                creds = Credentials.from_authorized_user_info(self._credentials)
            else:
                # Use service account or default credentials
                from google.auth import default
                creds, _ = default()
            
            self._service = build("docs", "v1", credentials=creds)
        return self._service
    
    def supports_mime_type(self, mime_type: str) -> bool:
        return mime_type == "application/vnd.google-apps.document"
    
    def supports_extension(self, extension: str) -> bool:
        return False  # Google Docs don't have extensions
    
    async def parse(self, doc_id: str) -> ParsedContent:
        service = self._get_service()
        
        # Get document
        doc = service.documents().get(documentId=doc_id).execute()
        
        content = doc.get("body", {}).get("content", [])
        text_parts = []
        headings = []
        tables = []
        images = []
        
        def extract_text(elements):
            for elem in elements:
                if "paragraph" in elem:
                    para = elem["paragraph"]
                    para_text = ""
                    for pe in para.get("elements", []):
                        if "textRun" in pe:
                            para_text += pe["textRun"].get("content", "")
                    
                    if para_text.strip():
                        text_parts.append(para_text)
                    
                    # Check for heading style
                    style = para.get("paragraphStyle", {})
                    named_style = style.get("namedStyleType", "")
                    if named_style.startswith("HEADING"):
                        level = int(named_style.replace("HEADING_", "")) if named_style != "HEADING" else 1
                        headings.append({
                            "level": level,
                            "text": para_text.strip(),
                            "position": len(text_parts) - 1
                        })
                
                elif "table" in elem:
                    table = elem["table"]
                    table_data = []
                    for row in table.get("tableRows", []):
                        row_data = []
                        for cell in row.get("tableCells", []):
                            cell_text = ""
                            for cell_elem in cell.get("content", []):
                                if "paragraph" in cell_elem:
                                    for pe in cell_elem["paragraph"].get("elements", []):
                                        if "textRun" in pe:
                                            cell_text += pe["textRun"].get("content", "")
                            row_data.append(cell_text.strip())
                        table_data.append(row_data)
                    tables.append({
                        "table_index": len(tables),
                        "data": table_data
                    })
                
                elif "tableOfContents" in elem:
                    # Handle TOC
                    pass
        
        extract_text(content)
        
        full_text = "\n\n".join(text_parts)
        
        return ParsedContent(
            text=full_text,
            structure={
                "headings": headings,
                "tables": tables,
                "images": images,
                "paragraphs": text_parts
            },
            metadata={
                "title": doc.get("title"),
                "document_id": doc_id,
                "revision_id": doc.get("revisionId"),
                "paragraph_count": len(text_parts),
                "table_count": len(tables)
            },
            source_type="google_docs"
        )


class DocumentParserFactory:
    def __init__(self):
        self.parsers: List[DocumentParser] = [
            TextParser(),
            PDFParser(),
            WordParser(),
        ]
        self._google_docs_parser: Optional[GoogleDocsParser] = None
    
    def set_google_credentials(self, credentials: Dict[str, Any]):
        self._google_docs_parser = GoogleDocsParser(credentials)
        self.parsers.append(self._google_docs_parser)
    
    def get_parser(self, mime_type: str = "", extension: str = "") -> Optional[DocumentParser]:
        for parser in self.parsers:
            if parser.supports_mime_type(mime_type) or parser.supports_extension(extension):
                return parser
        return None
    
    async def parse_bytes(self, content: bytes, mime_type: str = "", filename: str = "") -> ParsedContent:
        extension = ""
        if filename:
            extension = "." + filename.split(".")[-1] if "." in filename else ""
        
        parser = self.get_parser(mime_type, extension)
        if not parser:
            # Default to text parser
            parser = TextParser()
        
        return await parser.parse(content, filename)
    
    async def parse_google_doc(self, doc_id: str) -> ParsedContent:
        if not self._google_docs_parser:
            raise ValueError("Google Docs parser not configured. Call set_google_credentials first.")
        return await self._google_docs_parser.parse(doc_id)
