import base64
import html
import json
import zipfile
from datetime import datetime
from io import BytesIO
from typing import Any, Dict, List, Optional

from app.core.logging import logger
from app.services.ai_solution import (
    append_downloaded_assets_to_zip,
    extract_assignment_deliverables,
    extract_code_deliverables,
    extract_execution_metadata,
    strip_execution_metadata,
)
from app.services.terminal_renderer import TerminalCardRenderer


def build_solution_zip(
    assignment_id: int,
    raw_content: str,
) -> bytes:
    """Build a ZIP bundle containing generated source code files, terminal outputs, and assets."""
    clean_text = strip_execution_metadata(raw_content)
    deliverables = extract_assignment_deliverables(clean_text)
    files = deliverables.get("files", [])
    execution_results = extract_execution_metadata(raw_content)
    has_downloaded_assets = any(result.get("downloaded_assets") for result in execution_results)

    if not files and not has_downloaded_assets and not deliverables.get("tasks"):
        raise ValueError("No task deliverables were found in this solution.")

    archive = BytesIO()
    with zipfile.ZipFile(archive, mode="w", compression=zipfile.ZIP_DEFLATED) as bundle:
        for source_file in files:
            bundle.writestr(source_file["name"], source_file["content"])
        append_downloaded_assets_to_zip(bundle, execution_results, solution_markdown=clean_text, files=files)
        for index, output in enumerate(execution_results or deliverables.get("terminal_outputs", [])):
            filename = output.get("filename") or f"terminal_output_task_{output.get('task', index + 1)}.txt"
            text = output.get("terminal_output") or output.get("content") or "No terminal output was produced."
            bundle.writestr(f"terminal_output_{filename}.txt", text + "\n")

    return archive.getvalue()


def build_solution_html_report(
    assignment_title: str,
    course_name: str,
    instructor_name: str,
    student_name: str,
    student_id: str,
    due_at: Optional[datetime],
    raw_content: str,
    generated_images: Optional[List[Dict[str, Any]]] = None,
) -> str:
    """Build a standalone, printable HTML report of the generated solution with execution proof."""
    clean_content = strip_execution_metadata(raw_content)
    deliverables = extract_code_deliverables(clean_content)
    execution_results = extract_execution_metadata(raw_content)
    result_by_name = {item.get("filename"): item for item in execution_results}

    sections = []
    for source_file in deliverables.get("files", []):
        result = result_by_name.get(source_file["name"], {})
        output_text = result.get("terminal_output") or "No terminal output was produced."
        png = TerminalCardRenderer.generate_terminal_png(
            result.get("command") or f"Run {source_file['name']}",
            output_text,
            title=f"Assignment run — {source_file['name']}",
        )
        screenshot = f"data:image/png;base64,{base64.b64encode(png).decode('ascii')}"
        status = result.get("status", "not_run").replace("_", " ").title()
        sections.append(
            f'<section class="task"><h2>Task {html.escape(str(source_file["task"]))} · {html.escape(source_file["name"])}</h2>'
            f'<p class="meta">{html.escape(source_file["language"])} · Run status: {html.escape(status)} · Exit code: {html.escape(str(result.get("exit_code") if result.get("exit_code") is not None else "n/a"))}</p>'
            f'<h3>Source code</h3><pre><code>{html.escape(source_file["content"])}</code></pre>'
            f'<h3>Terminal output</h3><pre><code>{html.escape(output_text)}</code></pre>'
            f'<figure><img alt="Terminal screenshot for {html.escape(source_file["name"])}" src="{screenshot}"/><figcaption>Run evidence · {html.escape(source_file["name"])}</figcaption></figure></section>'
        )

    for generated_image in (generated_images or []):
        image_uri = generated_image.get("data_uri", "")
        prompt = html.escape(generated_image.get("prompt", "Generated assignment illustration"))
        if image_uri.startswith("data:image/"):
            sections.append(
                f'<section class="task"><h2>Generated illustration</h2>'
                f'<p class="meta">Prompt: {prompt}</p><figure><img alt="Assignment illustration" '
                f'src="{image_uri}"/><figcaption>Academic deliverable illustration</figcaption></figure></section>'
            )

    if not sections:
        sections.append(
            f'<section class="task"><h2>Solution</h2><pre><code>{html.escape(clean_content)}</code></pre><p class="meta">No separate code files were extracted from this solution.</p></section>'
        )

    due_str = due_at.strftime("%B %d, %Y") if due_at else "Not specified"
    now_str = datetime.now().strftime("%B %d, %Y")

    return f"""<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>{html.escape(assignment_title or 'Assignment')} · Solution report</title><style>
*{{box-sizing:border-box}}body{{margin:0;background:#f3f4f7;color:#20232d;font:15px/1.55 Arial,sans-serif}}.cover{{min-height:92vh;padding:12vh 12%;display:flex;flex-direction:column;justify-content:center;background:#fff;border-bottom:1px solid #ddd}}.brand{{color:#6557c9;font-weight:700;letter-spacing:.12em;text-transform:uppercase}}h1{{font-size:36px;line-height:1.2;margin:20px 0 30px}}.cover dl{{display:grid;grid-template-columns:180px 1fr;gap:10px 20px;max-width:650px}}dt{{color:#747987}}dd{{margin:0;font-weight:600}}main{{max-width:920px;margin:40px auto;padding:0 20px}}.task{{margin:24px 0;padding:30px;background:#fff;border:1px solid #e3e5eb;border-radius:12px;break-inside:avoid}}h2{{margin:0 0 6px;font-size:21px}}h3{{margin:24px 0 8px;font-size:13px;text-transform:uppercase;letter-spacing:.08em;color:#717687}}.meta,figcaption{{font-size:12px;color:#737888}}pre{{overflow:auto;white-space:pre-wrap;overflow-wrap:anywhere;margin:0;padding:16px;border-radius:8px;background:#111827;color:#ecf0f8;font:12px/1.6 Consolas,monospace}}figure{{margin:18px 0 0}}img{{display:block;width:100%;height:auto;border-radius:8px}}figcaption{{margin-top:7px}}@media print{{body{{background:white}}.cover{{min-height:95vh;page-break-after:always}}main{{margin:0 auto}}.task{{box-shadow:none}}}}
</style></head><body><header class="cover"><p class="brand">Uniflow · Academic report</p><h1>{html.escape(assignment_title or 'Assignment solution')}</h1><dl><dt>Course</dt><dd>{html.escape(course_name or 'Course')}</dd><dt>Student</dt><dd>{html.escape(student_name or 'Student')}</dd><dt>Student ID</dt><dd>{html.escape(str(student_id or 'Not provided'))}</dd><dt>Instructor</dt><dd>{html.escape(instructor_name or 'Instructor not provided')}</dd><dt>Due date</dt><dd>{html.escape(due_str)}</dd><dt>Report date</dt><dd>{now_str}</dd></dl></header><main>{''.join(sections)}</main></body></html>"""


def build_solution_docx_report(
    assignment_title: str,
    course_name: str,
    instructor_name: str,
    student_name: str,
    due_at: Optional[datetime],
    raw_content: str,
    base_docx_bytes: Optional[bytes] = None,
    generated_images: Optional[List[Dict[str, Any]]] = None,
) -> bytes:
    """Build or augment a Word .docx document with solution code, terminal outputs, and execution screenshots."""
    from docx import Document
    from docx.shared import Inches

    document = Document(BytesIO(base_docx_bytes)) if base_docx_bytes else Document()

    clean_content = strip_execution_metadata(raw_content)
    deliverables = extract_code_deliverables(clean_content)
    execution_results = extract_execution_metadata(raw_content)
    results_by_name = {item.get("filename"): item for item in execution_results}

    if not base_docx_bytes:
        document.add_heading("Uniflow · Academic report", 0)
        document.add_heading(assignment_title or "Assignment solution", level=1)
        document.add_paragraph(f"Course: {course_name or 'Course'}")
        document.add_paragraph(f"Student: {student_name or 'Student'}")
        document.add_paragraph(f"Instructor: {instructor_name or 'Not provided'}")
        document.add_paragraph(f"Due date: {due_at.strftime('%B %d, %Y') if due_at else 'Not specified'}")
        document.add_page_break()

        for source_file in deliverables.get("files", []):
            document.add_heading(f"Task {source_file['task']} · {source_file['name']}", level=1)
            document.add_heading("Source code", level=2)
            document.add_paragraph(source_file["content"])
            result = results_by_name.get(source_file["name"], {})
            document.add_heading("Terminal output", level=2)
            document.add_paragraph(result.get("terminal_output") or "No terminal output was produced.")

    if execution_results:
        document.add_page_break()
        document.add_heading("Code execution evidence", level=1)
        for result in execution_results:
            file_name = result.get("filename") or "Generated source"
            document.add_heading(file_name, level=2)
            run_state = result.get("status", "not_run").replace("_", " ").title()
            document.add_paragraph(f"Run status: {run_state} · Exit code: {result.get('exit_code') if result.get('exit_code') is not None else 'n/a'}")
            screenshot = TerminalCardRenderer.generate_terminal_png(
                result.get("command") or f"Run {file_name}",
                result.get("terminal_output") or "No terminal output was produced.",
                title=f"Assignment run — {file_name}",
            )
            document.add_picture(BytesIO(screenshot), width=Inches(6.2))
            document.add_paragraph(f"Run screenshot · {file_name}")

    if generated_images:
        document.add_page_break()
        document.add_heading("Generated illustrations", level=1)
        for image in generated_images:
            prompt = image.get("prompt", "")
            document.add_paragraph(f"Prompt: {prompt}")
            data_uri = image.get("data_uri", "")
            if data_uri.startswith("data:image/") and "," in data_uri:
                try:
                    document.add_picture(BytesIO(base64.b64decode(data_uri.split(",", 1)[1])), width=Inches(5.8))
                    document.add_paragraph("Figure: Academic deliverable illustration.")
                except (ValueError, OSError):
                    document.add_paragraph("Image data could not be embedded in this Word export.")

    output = BytesIO()
    document.save(output)
    return output.getvalue()
